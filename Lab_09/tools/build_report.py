"""Build Lab_09 report from measured benchmark artifacts."""
from __future__ import annotations

import json, os, subprocess
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS, ASSETS, OUT = ROOT/"benchmarks/results", ROOT/"report/assets", ROOT/"report/report.docx"
PYTHON = "/Users/denyszastavnyi/Documents/ChatGPT/проф пайтон/.venv/bin/python"

def fnt(n):
    p = "/System/Library/Fonts/Menlo.ttc"
    return ImageFont.truetype(p, n) if Path(p).exists() else ImageFont.load_default()

def terminal(text, path, title):
    lines = text.splitlines()[-22:]
    im = Image.new("RGB", (1600, 110+38*len(lines)), "#101820"); d = ImageDraw.Draw(im)
    d.rectangle((0,0,1600,66), fill="#263746"); d.text((35,16), title, font=fnt(28), fill="white")
    for i,line in enumerate(lines): d.text((28,82+i*38), line[:112], font=fnt(23), fill="#dce7f1")
    im.save(path)

def shade(cell, fill):
    x=OxmlElement("w:shd"); x.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(x)

def setup(d):
    s=d.sections[0]; s.top_margin=Cm(1.7); s.bottom_margin=Cm(1.6); s.left_margin=Cm(2.1); s.right_margin=Cm(1.7)
    n=d.styles["Normal"]; n.font.name="Times New Roman"; n.font.size=Pt(10.5); n.font.color.rgb=RGBColor(0,0,0); n.paragraph_format.space_after=Pt(4)
    for name,size in (("Title",22),("Heading 1",16),("Heading 2",13)):
        st=d.styles[name]; st.font.name="Arial"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    ppr=d.styles["Title"].element.get_or_add_pPr(); border=ppr.find(qn("w:pBdr"))
    if border is not None: ppr.remove(border)

def footer(d):
    p=d.sections[0].footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Заставний Денис • ФеП-32 • Lab_09 • ")
    x=OxmlElement("w:fldSimple"); x.set(qn("w:instr"),"PAGE"); p._p.append(x)

def table(d, headers, rows):
    t=d.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    flag=OxmlElement("w:tblHeader"); flag.set(qn("w:val"),"true"); t.rows[0]._tr.get_or_add_trPr().append(flag)
    for i,v in enumerate(headers):
        t.cell(0,i).text=v; shade(t.cell(0,i),"203864")
        for r in t.cell(0,i).paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for k,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); shade(cells[i], "F5F8FB" if k%2 else "FFFFFF")

def code(d, text):
    t=d.add_table(rows=1,cols=1); shade(t.cell(0,0),"F2F4F7"); r=t.cell(0,0).paragraphs[0].add_run(text.strip()); r.font.name="Menlo"; r.font.size=Pt(7.5)

def pic(d,path,caption,alt):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; s=p.add_run().add_picture(str(path),width=Cm(16))
    s._inline.docPr.set("title",caption); s._inline.docPr.set("descr",alt)
    p=d.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True

Q = [
("Що таке concurrency?","Організація кількох задач із перекриванням їх життєвих циклів."),("Що таке parallelism?","Фізичне одночасне виконання на кількох ядрах."),("Різниця concurrency і parallelism?","Перше описує структуру, друге — фактичну одночасність."),("Що таке process?","Ізольований екземпляр програми з власною пам'яттю."),("Що таке thread?","Потік виконання всередині процесу зі спільною пам'яттю."),("Process проти thread?","Process ізольований і дорожчий; threads ділять пам'ять."),("Що таке GIL?","Lock CPython, що допускає виконання Python bytecode одним thread."),("Роль GIL?","Спрощує керування об'єктами, але обмежує CPU threads."),("Що таке CPU-bound?","Задача, час якої визначають обчислення."),("Приклад CPU-bound?","Агрегація мільйона оцінок Python-циклами."),
("Що таке I/O-bound?","Задача, що переважно очікує файл, мережу або БД."),("Приклад I/O-bound?","Імпорт кількох CSV-файлів."),("Коли threading?","Для конкурентного I/O й операцій, що відпускають GIL."),("Коли multiprocessing?","Для достатньо великих незалежних CPU-bound задач."),("Що робить Thread?","Описує потік, який виконує callable."),("Що робить start()?","Запускає потік або процес."),("Що робить join()?","Очікує завершення потоку або процесу."),("Що таке race condition?","Залежність результату від порядку конкурентних доступів."),("Що таке critical section?","Код, що працює зі спільним mutable state."),("Для чого Lock?","Для взаємного виключення у critical section."),
("Що таке synchronization?","Координація доступу і порядку конкурентних задач."),("Вплив synchronization?","Гарантує correctness, але додає contention/overhead."),("Що таке ThreadPoolExecutor?","Пул потоків із submit/map та Future."),("Коли ThreadPoolExecutor?","Для багатьох однотипних I/O-bound операцій."),("Що таке ProcessPoolExecutor?","Пул процесів для паралельних callable."),("Коли ProcessPoolExecutor?","Для грубозернистих CPU-bound задач."),("Що робить executor.map()?","Розподіляє iterable і зберігає порядок результатів."),("Що таке worker?","Потік або процес, що виконує задачу пулу."),("Більше workers завжди краще?","Ні, overhead і contention можуть зрости."),("Що таке overhead?","Час запуску, планування, копіювання й координації."),
("Що таке serialization?","Кодування об'єктів для передачі між процесами."),("Чому multiprocessing повільний для малих задач?","IPC і запуск дорожчі за обчислення."),("Для чого timeit?","Для повторного мікробенчмарку."),("Для чого perf_counter()?","Для точного монотонного вимірювання інтервалів."),("Навіщо повтори benchmark?","Щоб зменшити шум і використати median."),("Що таке profiling?","Вимірювання витрат часу та пам'яті."),("Що робить cProfile?","Збирає статистику викликів Python-функцій."),("Що означає ncalls?","Кількість викликів."),("Що означає tottime?","Час функції без підвикликів."),("Що означає cumtime?","Час функції разом із підвикликами."),
("Що таке bottleneck?","Ділянка, що найбільше обмежує performance."),("Як знайти bottleneck?","Профілювати й сортувати за cumtime/tottime."),("Що робить tracemalloc?","Відстежує Python allocations і peak memory."),("Що таке peak memory?","Максимум пам'яті за час вимірювання."),("Що таке NumPy?","Бібліотека масивів і оптимізованих числових операцій."),("Що таке vectorization?","Заміна Python-циклів операціями над масивами."),("Чому NumPy швидший?","Цикли працюють в оптимізованому C-коді."),("Що таке caching?","Збереження результату для повторного використання."),("Що робить @cache?","Необмежено memoize-ить hashable аргументи."),("Що робить @lru_cache?","Memoization з LRU-витісненням."),
("Що таке cache hit?","Результат знайдено в кеші."),("Що таке cache miss?","Результат треба обчислити."),("Що таке invalidation?","Видалення застарілого кешу після зміни даних."),("Небезпека caching?","Stale data або зайве споживання пам'яті."),("Формула speedup?","T_baseline / T_optimized."),("Що означає speedup < 1?","Нова версія повільніша за baseline."),("Чому CPU threads можуть бути повільні?","GIL, scheduling і synchronization overhead."),("Коли ProcessPool прискорює CPU?","Коли chunks великі й перекривають IPC overhead."),("Як перевірити correctness?","Точні поля порівняти прямо, floats — з tolerance."),("Чому profile first?","Щоб оптимізувати реальний bottleneck, а не припущення."),]

def build():
    ASSETS.mkdir(parents=True,exist_ok=True)
    b=json.loads((RESULTS/"benchmark.json").read_text()); a=json.loads((RESULTS/"academic_benchmark.json").read_text()); io=json.loads((RESULTS/"io_benchmark.json").read_text()); c=json.loads((RESULTS/"cache_benchmark.json").read_text()); m=json.loads((RESULTS/"memory.json").read_text())
    env=os.environ.copy(); env["PYTHONPATH"]=str(ROOT/"src")
    r=subprocess.run([PYTHON,"-m","pytest","-q","--disable-warnings"],cwd=ROOT,env=env,text=True,capture_output=True,check=True)
    terminal(r.stdout+r.stderr,ASSETS/"pytest.png","pytest — Lab_09")
    lines=["dataset implementation workers time speedup"]+[f"{x['dataset_size']} {x['implementation']:<13} {x['workers']} {x['median_seconds']:.6f}s {x['speedup']:.2f}x" for x in b if x['dataset_size']==1_000_000]
    terminal("\n".join(lines),ASSETS/"benchmark.png","perf_counter — 1 000 000 grades")
    d=Document(); setup(d); footer(d)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n").bold=True; p.add_run("ЛЬВІВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ\nІМЕНІ ІВАНА ФРАНКА\nФакультет електроніки та комп'ютерних технологій")
    d.add_paragraph("\n"); p=d.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("ЗВІТ\nДО ЛАБОРАТОРНОЇ РОБОТИ №9")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Concurrency, parallelism, profiling та оптимізація\n").bold=True; p.add_run("Варіант 7 — система оцінювання студентів")
    d.add_paragraph("\n"); t=d.add_table(rows=2,cols=2); t.cell(0,0).text="Виконав:\nстудент групи ФеП-32\nЗаставний Денис"; t.cell(0,1).text="Перевірив:\nСвелеба Сергій Андрійович"; t.cell(1,0).merge(t.cell(1,1)); t.cell(1,0).text="Предмет: professional_python\nЛабораторна робота: Lab_09\nВаріант: 7"; shade(t.cell(1,0),"EAF1F8")
    d.add_paragraph("\n\n"); p=d.add_paragraph("Львів 2026"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; d.add_page_break()
    d.add_heading("1. Мета та вимоги",1); d.add_paragraph("Мета — знайти bottleneck, виміряти час і пам'ять, обґрунтовано застосувати concurrency/parallelism та оптимізувати систему оцінювання без зміни результатів.")
    table(d,["Вимога","Реалізація"],[["Набори","10 000; 100 000; 1 000 000 оцінок"],["CPU","Python loops, NumPy, ProcessPoolExecutor"],["I/O","8 CSV через ThreadPoolExecutor"],["Profiling","perf_counter, cProfile, tracemalloc"],["State/cache","Lock; hits/misses; invalidation"],["Workers","1, 2, 4; repeated median"],["Correctness","62 passed + equality checks"]])
    d.add_heading("2. Дані та baseline",1); d.add_paragraph("Детерміновані int32 масиви: student_id 0…9999, subject_id 0…19, grade 1…12, seed=7. Baseline Python-циклом рахує average/min/max/variance; окремо — середні за студентом і предметом, ranking та threshold.")
    code(d,"for raw in grades:\n    value = int(raw)\n    total += value\n    square_total += value * value\n    minimum = min(minimum, value)\n    maximum = max(maximum, value)")
    d.add_heading("3. Методика й bottleneck",1); d.add_paragraph("Кожний benchmark виконано 3 рази, взято median perf_counter(). cProfile показав домінування Python iteration та min/max. Speedup = Tbaseline / Toptimized; correctness перевірено перед збереженням результату."); d.add_page_break()
    d.add_heading("4. NumPy і основний експеримент",1); code(d,"student_counts = np.bincount(student_ids, minlength=students)\nstudent_totals = np.bincount(student_ids, weights=grades, minlength=students)\nranking = np.lexsort((np.arange(students), -student_averages))")
    rows=[]
    for n in (10_000,100_000,1_000_000):
        for impl,w in (("python_loops",1),("numpy",1),("process_pool",4)):
            x=next(z for z in b if z["dataset_size"]==n and z["implementation"]==impl and z["workers"]==w); rows.append([f"{n:,}".replace(","," "),impl,w,f"{x['median_seconds']:.6f}",f"{x['speedup']:.2f}×"])
    table(d,["Оцінок","Реалізація","Workers","Median,s","Speedup"],rows)
    d.add_paragraph("На 1 000 000 оцінок NumPy дав 71,64×. ProcessPool(4) — 0,73×, бо запуск процесів та передача chunks дорожчі за коротку агрегацію; це очікуваний негативний результат для дрібної гранулярності.")
    d.add_heading("5. Повна академічна аналітика",1); table(d,["Оцінок","Реалізація","Median,s","Speedup"],[[f"{x['dataset_size']:,}".replace(","," "),x['implementation'],f"{x['median_seconds']:.6f}",f"{x['speedup']:.2f}×"] for x in a]); d.add_paragraph("Для 1 000 000 рядків averages/dispersion/ranking/threshold скоротилися з 0,471917 s до 0,014976 s: 31,51×."); d.add_page_break()
    d.add_heading("6. Workers: ThreadPool і ProcessPool",1); wr=[]
    for impl in ("thread_pool","process_pool"):
        for w in (1,2,4):
            x=next(z for z in b if z['dataset_size']==1_000_000 and z['implementation']==impl and z['workers']==w); wr.append([impl,w,f"{x['median_seconds']:.6f}",f"{x['speedup']:.2f}×"])
    table(d,["Реалізація","Workers","Median,s","Speedup"],wr); d.add_paragraph("ThreadPool chunks викликають NumPy, який відпускає GIL; чистий Python CPU-loop так не масштабується. ProcessPool має IPC/startup overhead.")
    d.add_heading("7. I/O-bound CSV import",1); table(d,["Реалізація","Workers","Rows","Median,s","Speedup"],[[x['implementation'],x['workers'],x['rows'],f"{x['median_seconds']:.6f}",f"{x['speedup']:.2f}×"] for x in io]); d.add_paragraph("Для 8 локальних CSV/40 000 рядків SSD не створив достатнього очікування (≈1×), але ThreadPool коректно перекриває незалежні I/O операції.")
    d.add_heading("8. Synchronization",1); code(d,"with self._lock:\n    self._grades.append(value)\n# test: 500 concurrent add operations in 8 threads -> count == 500"); d.add_page_break()
    d.add_heading("9. Cache та invalidation",1); table(d,["Метрика","Значення"],[["Queries",c['repeated_queries']],["Hits / misses",f"{c['hits']} / {c['misses']}"],["Uncached",f"{c['uncached_seconds']:.6f} s"],["Cached",f"{c['cached_seconds']:.6f} s"],["Speedup",f"{c['speedup']:.2f}×"]]); d.add_paragraph("update_grade() завжди викликає invalidate(student_id); наступний get є miss і не повертає stale data.")
    d.add_heading("10. cProfile та tracemalloc",1); code(d,"\n".join((RESULTS/"cprofile.txt").read_text().splitlines()[:18])); table(d,["Implementation","Peak bytes"],[["Python baseline",m['baseline_peak_bytes']],["NumPy",m['numpy_peak_bytes']]]); d.add_paragraph("tracemalloc враховує Python allocations, але не весь native NumPy buffer, тому це діагностична метрика."); d.add_page_break()
    d.add_heading("11. Correctness і тести",1); d.add_paragraph("count/min/max/ranking/threshold порівняні точно; floats — np.isclose/allclose. Покрито генерацію, validation, pools, CSV, Lock, cache, profiling, benchmark, API, repository, DB і services."); pic(d,ASSETS/"pytest.png","Рисунок 1 — pytest","62 passed; один platform skip у sandbox"); d.add_paragraph("Єдиний skip спричинений забороною системних семафорів у test sandbox. Повний ProcessPool benchmark поза sandbox успішно виконав workers 1/2/4."); pic(d,ASSETS/"benchmark.png","Рисунок 2 — benchmark 1 000 000 оцінок","Реальні times і speedup"); d.add_page_break()
    d.add_heading("12. Аналіз та відтворення",1); d.add_paragraph("NumPy — найкращий вибір для щільної числової агрегації. ProcessPool корисний лише коли CPU chunks достатньо великі. ThreadPool доцільний для I/O. Cache найкраще працює для повторних незмінних запитів, але потребує invalidation. Lock забезпечує correctness shared mutable state."); code(d,"python -m pip install -e '.[test]'\npytest -q\npython tools/run_experiment.py"); d.add_paragraph("Артефакти збережено в benchmarks/results/: benchmark, academic, I/O, cache, cProfile і memory."); d.add_page_break()
    for part,items in enumerate((Q[:30],Q[30:])):
        d.add_heading(f"13.{part+1}. Контрольні питання {1 if part==0 else 31}–{30 if part==0 else 60}",1)
        for i,(q,ans) in enumerate(items,1+30*part): p=d.add_paragraph(); p.add_run(f"{i}. {q} ").bold=True; p.add_run(ans)
        d.add_page_break()
    d.add_heading("14. Висновок",1); d.add_paragraph("Для варіанта 7 реалізовано всі вимоги Lab_09: набір 1 000 000 оцінок, Python/NumPy/ProcessPool, ThreadPool для I/O, 1/2/4 workers, perf_counter, cProfile, tracemalloc, Lock, cache hits/misses та invalidation, repeated benchmark, speedup і correctness tests. NumPy прискорив загальну статистику у 71,64×, повну академічну аналітику — у 31,51×. Негативний speedup ProcessPool пояснено виміряним overhead."); d.add_paragraph("Репозиторій: https://github.com/DioDen5/student-grading-system-lab-09")
    OUT.parent.mkdir(parents=True,exist_ok=True); d.save(OUT); print(f"Saved {OUT}")

if __name__ == "__main__": build()
